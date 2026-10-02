from __future__ import annotations

import hashlib
import importlib.util
import json
import shutil
import sqlite3
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
SCHEMA_PATH = ROOT / "sql" / "task_contacts_schema.sql"
IMPORTER_PATH = ROOT / "scripts" / "import_confirmed_contacts.py"
MIGRATION_PATH = ROOT / "scripts" / "validate_task_contacts_migration.py"

HELD_TASK_IDS = (
    "H004",
    "F105",
    "F109",
    "F111",
    "F202",
    "F203",
    "F205",
    "F206",
    "F301",
    "F302",
    "F303",
    "F304",
)

CONFIRMED_IDS = (
    "A005",
    "A011",
    "F101",
    "M005",
    "M006",
    "R002",
    "R003",
    *(f"X{index:03d}" for index in range(1, 25)),
)
MULTIPLE_IDS = ("A012", "M007", *(f"X{index:03d}" for index in range(25, 32)))
CONDITIONAL_IDS = ("M009", *(f"X{index:03d}" for index in range(32, 42)))
TASK_IDS = CONFIRMED_IDS + MULTIPLE_IDS + CONDITIONAL_IDS + HELD_TASK_IDS


@dataclass(frozen=True)
class SyntheticContactDataset:
    root: Path
    db: Path
    backup: Path
    populated_db: Path
    workbook: Path
    tasks: Path
    schema: Path
    workbook_sha256: str
    db_sha256: str
    backup_sha256: str


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _status_for(task_id: str) -> str:
    if task_id in CONFIRMED_IDS:
        return "confirmed"
    if task_id in MULTIPLE_IDS:
        return "confirmed_multiple"
    if task_id in CONDITIONAL_IDS:
        return "conditional"
    held_index = HELD_TASK_IDS.index(task_id)
    if held_index < 4:
        return "ambiguous"
    if held_index == 4:
        return "conflict"
    if held_index < 7:
        return "probable_review"
    return "unmatched"


def _department_and_team(task_id: str) -> tuple[str, str]:
    if task_id in {"M005", "M007", "A005"}:
        return "건강증진과", "모자보건팀"
    if task_id == "R002":
        return "건강증진과", "지역보건팀"
    return "합성부서", "합성팀"


def _task_name(task_id: str) -> str:
    return f"합성 업무 {task_id}"


def _aliases_for(task_id: str, *, add_sixth: bool) -> list[dict[str, Any]]:
    meaningful = {
        "A011": "결핵",
        "A012": "결핵 검사",
        "F101": "결핵 안내",
        "R002": "어르신 오늘 건강",
        "R003": "연명치료",
        "H004": "금연아파트 지정",
    }.get(task_id, f"검색어 {task_id}")
    aliases = [
        {"text": meaningful, "weight": 10, "type": "synthetic"},
        *(
            {"text": f"합성 별칭 {task_id}-{index}", "weight": index, "type": "synthetic"}
            for index in range(1, 5)
        ),
    ]
    if add_sixth:
        aliases.append({"text": f"합성 별칭 {task_id}-5", "weight": 1, "type": "synthetic"})
    return aliases


def _build_tasks() -> list[dict[str, Any]]:
    tasks: list[dict[str, Any]] = []
    for index, task_id in enumerate(TASK_IDS):
        department, team = _department_and_team(task_id)
        tasks.append(
            {
                "id": task_id,
                "name": _task_name(task_id),
                "representative": f"합성 대표 {task_id}",
                "department": department,
                "team": team,
                "floor": "1층",
                "place": "합성실",
                "route": "합성 경로",
                "question": "합성 질문",
                "caution": "합성 주의",
                "script": "합성 안내",
                "priority": 100 if task_id == "R003" else 0,
                "status": "synthetic",
                "source": "synthetic fixture",
                "note": "",
                "locationCondition": "",
                "aliases": _aliases_for(task_id, add_sixth=index == 0),
            }
        )
    assert len(tasks) == 63
    assert sum(len(task["aliases"]) for task in tasks) == 316
    return tasks


def _phone(seed: int) -> str:
    return f"031-7000-{1000 + seed:04d}"


def _contact_plan() -> dict[str, dict[str, Any]]:
    plans: dict[str, dict[str, Any]] = {}
    phone_seed = 0

    for task_id in CONFIRMED_IDS:
        phone_seed += 1
        primary = _phone(phone_seed)
        if task_id in {"A011", "F101"}:
            primary = "031-5189-4364"
        elif task_id == "M005":
            primary = "031-5189-6944"
        elif task_id == "M006":
            primary = "031-5189-5023(기존 근거)"
        elif task_id == "R002":
            primary = "031-5189-5032"
        elif task_id == "A005":
            primary = "031-5189-5076"
        plans[task_id] = {"primary": primary, "additional": [], "roles": {}}

    for task_id in MULTIPLE_IDS:
        if task_id == "A012":
            plans[task_id] = {
                "primary": "031-5189-4364",
                "additional": [
                    "031-5189-4344",
                    "031-5189-4369",
                    "031-5189-4368",
                    "031-5189-4377",
                ],
                "roles": {
                    "031-5189-4344": "흉부 X선",
                    "031-5189-4369": "검체검사",
                    "031-5189-4368": "진단검사실",
                    "031-5189-4377": "민원접수",
                },
            }
            continue
        if task_id == "M007":
            plans[task_id] = {
                "primary": "031-5189-5076",
                "additional": ["031-5189-5075", "031-5189-4370", "031-5189-5085"],
                "roles": {},
            }
            continue
        phone_seed += 1
        primary = _phone(phone_seed)
        phone_seed += 1
        plans[task_id] = {"primary": primary, "additional": [_phone(phone_seed)], "roles": {}}

    for index, task_id in enumerate(CONDITIONAL_IDS):
        contact_count = 4 if index < 6 else 3
        values = []
        for _ in range(contact_count):
            phone_seed += 1
            values.append(_phone(phone_seed))
        if task_id == "M009":
            values[0] = "031-5189-5023(기존 근거)"
        has_primary = index < 5
        plans[task_id] = {
            "primary": values[0] if has_primary else "",
            "additional": values[1:] if has_primary else values,
            "roles": {},
        }

    for index, task_id in enumerate(HELD_TASK_IDS):
        status = _status_for(task_id)
        if status == "unmatched":
            plans[task_id] = {"primary": "", "additional": [], "roles": {}}
        else:
            phone_seed += 1
            plans[task_id] = {"primary": _phone(phone_seed), "additional": [], "roles": {}}
    return plans


def _create_baseline_db(path: Path, tasks: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(path) as connection:
        connection.execute("PRAGMA foreign_keys = ON")
        connection.executescript(
            """
            CREATE TABLE tasks (
                id TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                representative TEXT,
                department TEXT,
                team TEXT,
                floor TEXT,
                place TEXT,
                route TEXT,
                question TEXT,
                caution TEXT,
                script TEXT,
                priority INTEGER NOT NULL DEFAULT 0,
                status TEXT,
                source TEXT,
                note TEXT,
                location_condition TEXT,
                contact_name TEXT,
                contact_role TEXT,
                phone TEXT,
                contact_verified_at TEXT,
                active INTEGER NOT NULL DEFAULT 1,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            CREATE TABLE aliases (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id TEXT NOT NULL REFERENCES tasks(id) ON DELETE CASCADE,
                text TEXT NOT NULL,
                weight INTEGER NOT NULL DEFAULT 1,
                type TEXT,
                UNIQUE(task_id, text)
            );
            CREATE TABLE event_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                event_type TEXT NOT NULL,
                task_id TEXT REFERENCES tasks(id) ON DELETE SET NULL,
                result_count INTEGER,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            """
        )
        for task in tasks:
            connection.execute(
                """
                INSERT INTO tasks(
                    id, name, representative, department, team, floor, place, route,
                    question, caution, script, priority, status, source, note,
                    location_condition, active
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1)
                """,
                (
                    task["id"],
                    task["name"],
                    task["representative"],
                    task["department"],
                    task["team"],
                    task["floor"],
                    task["place"],
                    task["route"],
                    task["question"],
                    task["caution"],
                    task["script"],
                    task["priority"],
                    task["status"],
                    task["source"],
                    task["note"],
                    task.get("locationCondition")
                    or task.get("location_condition")
                    or "",
                ),
            )
            connection.executemany(
                "INSERT INTO aliases(task_id, text, weight, type) VALUES (?, ?, ?, ?)",
                (
                    (task["id"], alias["text"], alias["weight"], alias["type"])
                    for alias in task["aliases"]
                ),
            )
        connection.executemany(
            "INSERT INTO event_logs(event_type, result_count) VALUES ('synthetic', ?)",
            ((index,) for index in range(114)),
        )


def _create_workbook(path: Path, tasks: list[dict[str, Any]]) -> None:
    final_headers = (
        "업무ID",
        "검색 업무명",
        "최종 부서",
        "최종 팀",
        "최종 상태",
        "최종 대표전화",
        "최종 추가연락처",
        "PDF 대조판정",
        "PDF 확인 내선",
        "PDF 근거 업무",
        "PDF 근거 위치",
        "기존 상태",
        "기존 대표전화",
        "검색엔진 반영 권고",
        "확정일",
    )
    supplement_headers = (
        "업무ID",
        "대표 연락처 역할",
        "적용 조건",
        "공식 출처 URL",
        "공식 매칭 근거",
    )
    evidence_headers = ("정규화 연락처", "분장사무 요약", "예외·비고", "출처 문서")
    plans = _contact_plan()
    workbook = Workbook()
    final_sheet = workbook.active
    final_sheet.title = "PDF대조_확정_63"
    final_sheet.append(final_headers)
    supplement_sheet = workbook.create_sheet("업무별_연락처_63")
    supplement_sheet.append(supplement_headers)
    evidence_sheet = workbook.create_sheet("PDF내선_근거_48")
    evidence_sheet.append(evidence_headers)
    admin_sheet = workbook.create_sheet("보건행정과_내선_25")
    admin_sheet.append(evidence_headers)

    for task in tasks:
        task_id = task["id"]
        status = _status_for(task_id)
        plan = plans[task_id]
        final_sheet.append(
            (
                task_id,
                task["name"],
                task["department"],
                task["team"],
                status,
                plan["primary"],
                "; ".join(plan["additional"]),
                "합성 대조",
                "합성 내선",
                "합성 업무 근거",
                "합성 위치 근거",
                "합성 기존 상태",
                "",
                "합성 반영 권고",
                "2026-08-28",
            )
        )
        base_role = "결핵 문의" if task_id in {"A011", "F101"} else "합성 문의"
        supplement_sheet.append(
            (
                task_id,
                base_role,
                "합성 조건" if status == "conditional" else "",
                f"https://www.hscity.go.kr/health/synthetic/{task_id}",
                "합성 매칭 근거",
            )
        )
        for raw_phone in [plan["primary"], *plan["additional"]]:
            phone = str(raw_phone).split("(", 1)[0]
            if not phone:
                continue
            role = plan["roles"].get(phone, f"합성 용도 {phone[-4:]}")
            evidence_sheet.append((phone, role, "", "합성 문서"))
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def build_synthetic_contact_dataset(root: Path) -> SyntheticContactDataset:
    tasks = _build_tasks()
    db = root / "data" / "health_search.db"
    backup = root / "data" / "backups" / "synthetic-backup.db"
    populated_db = root / "data" / "health_search.with-contacts.db"
    workbook = root / "data" / "source" / "synthetic-contacts.xlsx"
    tasks_path = root / "data" / "tasks.json"
    tasks_path.parent.mkdir(parents=True, exist_ok=True)
    tasks_path.write_text(json.dumps(tasks, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    _create_baseline_db(db, tasks)
    _create_workbook(workbook, tasks)

    backup.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(f"file:{db.resolve().as_posix()}?mode=ro", uri=True) as source:
        with sqlite3.connect(backup) as destination:
            source.backup(destination)
    shutil.copy2(db, populated_db)

    workbook_hash = file_sha256(workbook)
    importer = _load_module("synthetic_contact_importer", IMPORTER_PATH)
    importer.EXPECTED_WORKBOOK_SHA256 = workbook_hash
    payload, report = importer.run_import(
        workbook,
        tasks_path,
        db,
        expected_sha256=workbook_hash,
    )
    assert report["validation_passed"], report["errors"]
    migration = _load_module("synthetic_contact_migration", MIGRATION_PATH)
    rows = migration.build_contact_rows(payload, importer)
    with sqlite3.connect(populated_db, isolation_level=None) as connection:
        connection.execute("PRAGMA foreign_keys = ON")
        connection.execute("BEGIN IMMEDIATE")
        migration.migrate_schema(connection, SCHEMA_PATH)
        migration.apply_contact_rows(connection, rows, importer)
        connection.commit()
        validation = migration.validate_database(connection)
        held_rows = connection.execute(
            f"SELECT COUNT(*) FROM task_contacts WHERE active = 1 AND task_id IN "
            f"({','.join('?' for _ in HELD_TASK_IDS)})",
            HELD_TASK_IDS,
        ).fetchone()[0]
    assert validation["total_contact_rows"] == 93
    assert validation["staging_task_count"] == 51
    assert validation["primary_contact_rows"] == 45
    assert validation["tasks_with_multiple_contact_rows"] == 20
    assert validation["duplicate_count"] == 0
    assert validation["primary_duplicate_count"] == 0
    assert validation["orphan_task_id_count"] == 0
    assert validation["foreign_key_check"] == []
    assert validation["integrity_check"] == ["ok"]
    assert held_rows == 0

    return SyntheticContactDataset(
        root=root,
        db=db,
        backup=backup,
        populated_db=populated_db,
        workbook=workbook,
        tasks=tasks_path,
        schema=SCHEMA_PATH,
        workbook_sha256=workbook_hash,
        db_sha256=file_sha256(db),
        backup_sha256=file_sha256(backup),
    )


@pytest.fixture(scope="session")
def synthetic_contact_dataset(tmp_path_factory) -> SyntheticContactDataset:
    return build_synthetic_contact_dataset(tmp_path_factory.mktemp("synthetic-contact-data"))


PUBLIC_GUIDANCE_CONTACTS = {
    'H003': (('031-5189-6923', "대표전화", '비만타파·근력강화 운동교실', True),),
    'M001': (('031-5189-6944', "대표전화", '희귀질환자 의료비 지원 신청', True),),
    'R001': (('031-5189-5058', "대표전화", '암환자 의료비 지원 신청', True),),
    'R003': (('031-5189-5058', "대표전화", '사전연명의료의향서 등록', True),),
    'R004': (('031-5189-5058', "대표전화", '장기기증희망 등록', True),),
    'R005': (('031-5189-4342', "대표전화", '무릎인공관절 수술비 지원 신청', True),),
    'R006': (('031-5189-4342', "대표전화", '휠체어 대여·반납', True),),
    'D002': (('031-5189-4719', "대표전화", '치매쉼터·치매예방교실', True),),
    'A008': (('031-5189-5094', "대표전화", '성매개감염병 상담·관리', True),),
    'A017': (('031-5189-4377', "대표전화", '의약무 관련 제증명 수납', True),),
    'A018': (('031-5189-4377', "대표전화", '예방접종 수납', True),),
    'F107': (('031-5189-4719', "대표전화", '치매안심쉼터 위치 안내', True),),
    'F112': (('031-352-0175', "대표전화", '정신건강복지센터 위치 안내', True),),
    'F305': (('031-5189-6916', "대표전화", '건강증진과 과장실 위치 안내', True),),
    "A002": (
        ("031-5189-4369", "대표전화", "진단검사실", True),
        ("031-5189-4368", "추가 연락처", "진단검사실 운영", False),
    ),
    "A005": (
        ("031-5189-5076", "대표전화", "임산부 등록·일반 안내", True),
        ("031-5189-4370", "추가 연락처", "모자보건 상담·신청", False),
        ("031-5189-5085", "추가 연락처", "모자보건 상담·신청", False),
    ),
    "A009": (
        ("031-5189-5094", "대표전화", "성매개감염병 상담·관리", True),
        ("031-5189-4369", "검체검사", "검체검사", False),
        ("031-5189-4368", "진단검사실", "진단검사실", False),
    ),
    "M009": (
        ("031-5189-4375", "대표전화", "예방접종실 일반 문의", True),
        ("031-5189-4376", "추가 연락처", "위탁의료기관 예방접종", False),
    ),
    "F110": (("031-5189-5076", "대표전화", "모자보건교육", True),),
    "A004": (("031-5189-5094", "대표전화", "HIV 익명검사·상담", True),),
    "F102": (("031-5189-4344", "대표전화", "영상의학실", True),),
    "F104": (
        ("031-5189-4369", "대표전화", "진단검사실", True),
        ("031-5189-4368", "추가 연락처", "진단검사실 운영", False),
    ),
    "F106": (
        ("031-5189-4342", "대표전화", "재활보건실·보조기기", True),
        ("031-5189-5058", "추가 연락처", "재활사업", False),
    ),
    "F204": (("031-5189-4371", "대표전화", "금연상담실", True),),
    "A001": (("031-5189-4378", "대표전화", "진료실", True),),
    "A003": (("031-5189-4377", "대표전화", "제증명 발급", True),),
    "A007": (("031-5189-5093", "대표전화", "방역·소독 업무", True),),
    "A011": (("031-5189-4364", "대표전화", "결핵 상담·관리", True),),
    "A012": (
        ("031-5189-4364", "대표전화", "결핵 검사 안내", True),
        ("031-5189-4344", "흉부 X선", "흉부 X선", False),
        ("031-5189-4369", "검체검사", "검체검사", False),
        ("031-5189-4368", "진단검사실", "진단검사실", False),
        ("031-5189-4377", "민원접수", "민원접수", False),
    ),
    "A019": (("031-5189-4344", "대표전화", "영상의학실 골다공증 검사", True),),
    "F101": (("031-5189-4364", "대표전화", "결핵 안내", True),),
    "F103": (("031-5189-4378", "대표전화", "진료실", True),),
    "F108": (("031-5189-4377", "대표전화", "민원실", True),),
    "F201": (("031-5189-4374", "대표전화", "만성질환관리센터", True),),
    "H001": (("031-5189-4371", "대표전화", "금연클리닉", True),),
    "H002": (("031-5189-4374", "대표전화", "동탄 만성질환관리센터", True),),
    "M002": (("031-5189-6944", "대표전화", "선천성대사이상 지원", True),),
    "M003": (
        ("031-5189-6943", "대표전화", "난임 지원 사업 담당", True),
        ("031-5189-4370", "모자보건 지원 사업 상담·신청 접수", "모자보건 상담·신청", False),
        ("031-5189-5085", "모자보건 지원 사업 상담·신청 접수", "모자보건 상담·신청", False),
    ),
    "M004": (
        ("031-5189-4370", "대표전화", "산모·신생아 지원 신청", True),
        ("031-5189-5085", "모자보건 지원 사업 상담·신청 접수", "모자보건 상담·신청", False),
        ("031-5189-6944", "영유아·임산부 지원", "영유아·임산부 지원", False),
    ),
    "M005": (
        ("031-5189-6944", "대표전화", "고위험 임산부 의료비", True),
        ("031-5189-4370", "모자보건 지원 사업 상담·신청 접수", "모자보건 상담·신청", False),
        ("031-5189-5085", "모자보건 지원 사업 상담·신청 접수", "모자보건 상담·신청", False),
    ),
    "M006": (
        ("031-5189-6944", "대표전화", "기저귀·조제분유 지원", True),
        ("031-5189-4370", "모자보건 지원 사업 상담·신청 접수", "모자보건 상담·신청", False),
        ("031-5189-5085", "모자보건 지원 사업 상담·신청 접수", "모자보건 상담·신청", False),
        ("031-5189-5023", "추가 연락처", "기저귀 지원", False),
    ),
    "M008": (("031-5189-6944", "대표전화", "미숙아·선천성이상아 의료비", True),),
    "R002": (
        ("031-5189-5032", "대표전화", "어르신 건강관리 문의", True),
        ("031-5189-4778", "권역별 방문건강 문의", "방문건강 문의", False),
        ("031-5189-4779", "권역별 방문건강 문의", "방문건강 문의", False),
        ("031-5189-6933", "권역별 방문건강 문의", "방문건강 문의", False),
        ("031-5189-6946", "권역별 방문건강 문의", "방문건강 문의", False),
    ),
}


@pytest.fixture(scope="session")
def public_guidance_contact_db(tmp_path_factory) -> Path:
    """Build the public API fixture without the operating DB or source workbook."""

    root = tmp_path_factory.mktemp("public-guidance-api-data")
    database = root / "health_search.db"
    tasks = json.loads((ROOT / "data" / "tasks.json").read_text(encoding="utf-8"))
    _create_baseline_db(database, tasks)

    with sqlite3.connect(database) as connection:
        connection.execute("PRAGMA foreign_keys = ON")
        connection.executescript(SCHEMA_PATH.read_text(encoding="utf-8"))
        source_row = 2
        for task_id, contacts in PUBLIC_GUIDANCE_CONTACTS.items():
            status = "confirmed_multiple" if len(contacts) > 1 else "confirmed"
            for phone, label, role, is_primary in contacts:
                connection.execute(
                    """
                    INSERT INTO task_contacts(
                        task_id, phone, display_phone, label, note, contact_role,
                        condition_text, is_primary, source_url, source_urls_json,
                        source_sheet, source_row, verified_at, match_status, active
                    ) VALUES (?, ?, ?, ?, NULL, ?, NULL, ?, ?, '[]', ?, ?, ?, ?, 1)
                    """,
                    (
                        task_id,
                        phone,
                        phone,
                        label,
                        role,
                        int(is_primary),
                        "https://www.hscity.go.kr/health/",
                        "synthetic-public-guidance",
                        source_row,
                        "2026-08-28",
                        status,
                    ),
                )
                source_row += 1
        assert connection.execute("PRAGMA foreign_key_check").fetchall() == []
        assert connection.execute("PRAGMA integrity_check").fetchone() == ("ok",)
    return database
